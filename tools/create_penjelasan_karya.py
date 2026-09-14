from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "docs"
OUTPUT_PATH = OUTPUT_DIR / "Penjelasan Karya Ibuki Study.docx"

BLACK = "000000"
PURPLE = "6F1D9B"
DEEP_PURPLE = "42105F"
LIGHT_PURPLE = "F5EFF9"
PALE_PURPLE = "FAF7FC"
LIGHT_GRAY = "D9D9D9"
MID_GRAY = "666666"
PALE_GRAY = "F7F7F7"
WHITE = "FFFFFF"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Aptos", size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, size=9, color=MID_GRAY)


def add_kicker(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text.upper())
    set_run_font(r, size=9, bold=True, color=PURPLE)
    r.font.all_caps = True
    r.font.letter_spacing = Pt(1.2)
    return p


def add_body(doc, text, bold_lead=None, keep=False):
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    set_run_font(r, size=10.7)
    return p


def add_numbered(doc, title, body):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    r1 = p.add_run(title + ". ")
    set_run_font(r1, size=10.8, bold=True)
    r2 = p.add_run(body)
    set_run_font(r2, size=10.8)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_fill(cell, DEEP_PURPLE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=150, bottom=150)
        if widths:
            set_cell_width(cell, widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=9.5, bold=True, color=WHITE)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci, value in enumerate(row):
            cell = cells[ci]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                set_cell_width(cell, widths[ci])
            if ri % 2 == 1:
                set_cell_fill(cell, PALE_PURPLE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_run_font(r, size=9.2)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_section_title(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=19, bold=True, color=BLACK)
    return p


def add_subheading(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=13, bold=True, color=BLACK)
    return p


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = RGBColor.from_string(BLACK)

    body = styles["Body Text"]
    body.font.name = "Aptos"
    body._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    body._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    body.font.size = Pt(10.8)
    body.paragraph_format.space_after = Pt(7)
    body.paragraph_format.line_spacing = 1.13

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(34)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(10)

    for style_name, size in (("Heading 1", 19), ("Heading 2", 13)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(14 if style_name == "Heading 1" else 9)
        style.paragraph_format.space_after = Pt(6)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Ibuki Study  |  Penjelasan Karya  |  ")
    set_run_font(r, size=9, color=MID_GRAY)
    add_page_number(p)


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    # Cover
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(52)
    add_kicker(doc, "Ibuki Study")
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Penjelasan Karya Ibuki Study")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(26)
    r = subtitle.add_run("Ide  Cara Kerja  Fitur Unik")
    set_run_font(r, size=16, color=DEEP_PURPLE, bold=True)

    opening = doc.add_paragraph()
    opening.paragraph_format.space_after = Pt(18)
    opening.paragraph_format.line_spacing = 1.18
    r = opening.add_run(
        "Ibuki Study adalah ruang belajar berbasis AI yang mengubah catatan mentah menjadi pengetahuan terstruktur, "
        "jawaban yang dapat ditelusuri ke sumber, dan latihan aktif. Karya ini dirancang agar pelajar tidak berhenti "
        "pada rangkuman otomatis, tetapi dapat memeriksa dasar jawaban, menghubungkan materi, dan menguji pemahaman."
    )
    set_run_font(r, size=13, color=BLACK)

    add_table(
        doc,
        ["Dokumen", "Cakupan", "Pengguna"],
        [["Acuan penilaian karya", "Seluruh aplikasi web dan Android", "Juri dan tim demonstrasi"]],
        widths=[2.0, 2.45, 2.35],
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("Kesimpulan utama")
    set_run_font(r, size=11, bold=True, color=PURPLE)
    add_body(
        doc,
        "Pembeda utama Ibuki Study adalah keterhubungan antara sumber, penjelasan AI, organisasi pengetahuan, "
        "dan aktivitas belajar. Satu materi dapat ditelusuri ke bukti asalnya, dipakai sebagai konteks percakapan, "
        "dihubungkan dengan materi lain, lalu diubah menjadi kuis, flashcard, atau panduan belajar yang dapat dibagikan.",
    )

    page_break(doc)

    # Summary and idea
    add_section_title(doc, "Ringkasan Karya")
    add_body(
        doc,
        "Pelajar sering menyimpan foto papan tulis, PDF, dan catatan teks di tempat yang berbeda. Informasi tersebut "
        "sulit dicari kembali, hubungan antar topik tidak terlihat, dan chatbot umum dapat menjawab tanpa menunjukkan "
        "asal informasi. Ibuki Study menyatukan proses pengumpulan, pemahaman, verifikasi, pengorganisasian, dan latihan "
        "dalam satu alur kerja.",
    )
    add_table(
        doc,
        ["Kebutuhan Pelajar", "Respons Ibuki Study", "Hasil yang Diharapkan"],
        [
            ["Catatan tersebar dalam beberapa format", "Unggah gambar, teks, dan PDF ke satu notebook", "Materi lebih mudah ditemukan dan dipelajari ulang"],
            ["Rangkuman AI sulit dipercaya", "Sitasi per klaim dengan tampilan sumber", "Pelajar dapat memeriksa jawaban secara langsung"],
            ["Belajar pasif dari rangkuman", "Kuis, flashcard, dan panduan belajar", "Pemahaman diuji melalui latihan aktif"],
            ["Hubungan antar topik tidak terlihat", "Pohon materi dan graf relasi", "Struktur serta keterkaitan pengetahuan lebih jelas"],
        ],
        widths=[2.05, 2.55, 2.2],
    )

    add_section_title(doc, "Ide Karya")
    add_subheading(doc, "Masalah yang Diangkat")
    add_body(
        doc,
        "Banyak alat belajar AI berfokus pada pembuatan teks baru. Masalah pelajar sebenarnya lebih luas: sumber perlu "
        "diubah menjadi catatan yang rapi, isi perlu dapat diverifikasi, dan hasilnya harus mendorong latihan. Tanpa "
        "keterhubungan ini, AI hanya menjadi mesin jawaban dan tidak membentuk sistem belajar yang berkelanjutan.",
    )
    add_subheading(doc, "Gagasan Solusi")
    add_body(
        doc,
        "Kami merancang Ibuki Study sebagai basis pengetahuan pribadi. AI membantu membaca dan menyusun materi, tetapi "
        "pengguna tetap mengendalikan sumber yang dipakai, struktur materi, serta keputusan untuk menyimpan hasil. "
        "Verifikasi sumber ditempatkan di dalam pengalaman belajar agar pengguna dapat berpindah dari jawaban ke bukti "
        "tanpa keluar dari aplikasi.",
    )
    add_subheading(doc, "Sasaran Pengguna")
    add_body(
        doc,
        "Sasaran utama adalah pelajar yang belajar dari campuran foto catatan, papan tulis, berkas teks, dan dokumen. "
        "Aplikasi juga relevan bagi tim belajar yang ingin menata materi per mata pelajaran dan memakai AI dengan konteks "
        "yang dapat dipilih secara eksplisit.",
    )

    page_break(doc)

    # How it works
    add_section_title(doc, "Cara Kerja")
    add_body(
        doc,
        "Alur kerja utama bergerak dari sumber mentah menuju aktivitas belajar. Setiap tahap menyimpan hasil yang dapat "
        "dipakai kembali oleh tahap berikutnya.",
    )
    add_numbered(doc, "Masuk dan memuat profil", "Pengguna masuk dengan email dan kata sandi atau Google. Supabase mengelola sesi, profil, dan pemulihan kata sandi untuk web maupun Android.")
    add_numbered(doc, "Mengumpulkan sumber", "Pengguna memasukkan gambar, berkas teks, atau PDF. Gambar diperkecil sebelum dikirim, teks dibaca langsung, dan PDF diekstrak per halaman menggunakan PDF.js.")
    add_numbered(doc, "Menganalisis dan menyusun notebook", "Model visi atau model percakapan menghasilkan judul, transkripsi, rangkuman Markdown, mata pelajaran, serta jalur materi. Instruksi model meminta sitasi granular pada klaim yang berasal dari sumber.")
    add_numbered(doc, "Meninjau sebelum menyimpan", "Judul, mata pelajaran, topik, rangkuman, dan transkripsi dapat diperiksa serta diubah. Berkas asli diunggah ke penyimpanan dan metadata notebook disimpan ke basis data.")
    add_numbered(doc, "Mengatur basis pengetahuan", "Notebook ditempatkan pada pohon materi dengan kedalaman fleksibel. Pengguna dapat menambah, mengganti nama, memindahkan, mengelompokkan, atau meminta usulan pengelompokan AI.")
    add_numbered(doc, "Bertanya dengan konteks terpilih", "Chat tersedia pada tingkat notebook, mata pelajaran, dan seluruh koleksi. Panel sumber memungkinkan pengguna memilih cabang atau notebook yang masuk ke konteks. Jawaban diberikan secara streaming dan riwayat disimpan per cakupan.")
    add_numbered(doc, "Memeriksa sitasi", "Penanda sitasi pada rangkuman dan chat membuka sumber terkait. Teks dicari dan disorot secara toleran terhadap variasi spasi, PDF menampilkan penanda halaman, dan gambar dapat menyorot kotak koordinat yang dirujuk.")
    add_numbered(doc, "Berlatih dan menjaga progres", "Konteks notebook atau mata pelajaran dapat diubah menjadi kuis pilihan ganda dan flashcard. Kuis menyimpan jawaban, status, nilai terbaik, serta hari belajar. Data streak disinkronkan ke widget Android.")
    add_numbered(doc, "Membuat keluaran baru", "Dari percakapan, pengguna dapat menghasilkan panduan belajar untuk diekspor sebagai PDF atau materi baru yang ditinjau sebelum menjadi notebook.")

    add_subheading(doc, "Aliran Data Utama")
    add_table(
        doc,
        ["Tahap", "Masukan", "Proses", "Keluaran"],
        [
            ["Akuisisi", "Gambar, TXT, PDF", "Kompresi atau ekstraksi", "Daftar sumber"],
            ["Pemahaman", "Isi sumber", "OCR, sintesis, klasifikasi", "Notebook terstruktur"],
            ["Konteks", "Notebook pilihan", "Perakitan konteks dan registri sumber", "Prompt berbasis materi"],
            ["Interaksi", "Pertanyaan atau gambar", "Respons streaming dan tool action", "Jawaban, kuis, flashcard, panduan"],
            ["Penyimpanan", "Notebook dan progres", "Supabase Database dan Storage", "Riwayat lintas sesi"],
        ],
        widths=[1.05, 1.55, 2.25, 1.95],
    )

    page_break(doc)

    # Unique features
    add_section_title(doc, "Fitur Unik")
    add_subheading(doc, "Sitasi yang Dapat Dibuka Kembali")
    add_body(
        doc,
        "Setiap sumber diberi nomor konsisten dan nomor tersebut dipakai oleh prompt AI serta renderer antarmuka. "
        "Sitasi tidak berhenti sebagai catatan kaki. Pengguna dapat membuka kutipan yang dirujuk, menuju bagian teks "
        "yang disorot, halaman PDF, atau area spesifik pada foto catatan. Mekanisme ini membuat proses pemeriksaan menjadi "
        "bagian dari pengalaman utama.",
    )
    add_subheading(doc, "Konteks Chat yang Dikendalikan Pengguna")
    add_body(
        doc,
        "Pengguna memilih sendiri notebook yang boleh menjadi konteks. Pilihan dapat dilakukan pada satu dokumen, satu "
        "cabang materi, satu mata pelajaran, atau seluruh koleksi. Aplikasi menyimpan snapshot registri sumber pada setiap "
        "jawaban AI sehingga nomor sitasi lama tetap mengarah ke sumber yang benar walaupun pilihan konteks berubah.",
    )
    add_subheading(doc, "Dua Cara Melihat Pengetahuan")
    add_body(
        doc,
        "Pohon materi menunjukkan hierarki dari mata pelajaran menuju submateri, sedangkan graf menunjukkan hubungan "
        "semantik antar notebook. Relasi yang didukung adalah prasyarat, pengembangan dari materi sebelumnya, keterkaitan, "
        "perbandingan, dan contoh. Hubungan dapat dibuat manual atau disarankan AI; penyegaran AI tidak menimpa hubungan manual.",
    )
    add_subheading(doc, "Siklus Belajar Aktif")
    add_body(
        doc,
        "Kuis dapat diatur berdasarkan tingkat kesulitan, jumlah soal, dan fokus topik. Percobaan yang belum selesai dapat "
        "dilanjutkan, nilai terbaik disimpan, dan setiap jawaban mencatat hari belajar. Flashcard menyediakan mode balik "
        "pertanyaan dan jawaban, sedangkan widget Android mengingatkan status belajar harian.",
    )
    add_subheading(doc, "Percakapan Menjadi Artefak Belajar")
    add_body(
        doc,
        "Pesan atau gambar dalam chat dapat diubah menjadi panduan belajar yang diekspor sebagai PDF atau menjadi notebook "
        "baru. Untuk notebook baru, pengguna melihat pratinjau judul, mata pelajaran, topik, rangkuman, dan transkripsi "
        "sebelum menyimpan. Alur ini menjaga keputusan akhir pada pengguna.",
    )

    add_table(
        doc,
        ["Fitur", "Nilai bagi Juri", "Bukti Implementasi"],
        [
            ["Sitasi interaktif", "Transparansi dan verifikasi jawaban", "Registri sumber, kutipan, halaman, dan bbox"],
            ["Pemilih konteks", "Kontrol pengguna atas dasar jawaban", "Seleksi pohon dan snapshot per pesan"],
            ["Pohon dan graf", "Organisasi struktural sekaligus semantik", "Hierarki tanpa batas dan lima tipe relasi"],
            ["Kuis berprogres", "Menghubungkan pengetahuan dengan latihan", "Resume, nilai terbaik, streak, widget"],
            ["Chat ke artefak", "Hasil AI dapat dipakai kembali", "PDF panduan dan notebook hasil percakapan"],
        ],
        widths=[1.45, 2.5, 2.85],
    )

    page_break(doc)

    # Architecture
    add_section_title(doc, "Arsitektur dan Teknologi")
    add_body(
        doc,
        "Aplikasi memakai arsitektur klien modern dengan layanan cloud untuk autentikasi, data, dan berkas. Pemrosesan AI "
        "mengikuti antarmuka chat completions yang kompatibel dengan OpenAI, sedangkan Capacitor membungkus pengalaman web "
        "menjadi aplikasi Android.",
    )
    add_table(
        doc,
        ["Lapisan", "Teknologi", "Peran"],
        [
            ["Antarmuka", "Vue 3, Vite, CSS", "Dashboard, notebook, chat, kuis, graf, pengaturan"],
            ["AI", "API chat completions kompatibel OpenAI", "Analisis sumber, chat streaming, kuis, flashcard, pengelompokan, relasi"],
            ["Dokumen", "PDF.js, KaTeX, html2pdf.js", "Ekstraksi PDF, rumus matematika, ekspor panduan"],
            ["Visualisasi", "D3 Force", "Tata letak graf hubungan materi yang interaktif"],
            ["Backend", "Supabase Auth, Database, Storage", "Akun, data pribadi pengguna, sumber, riwayat, progres"],
            ["Mobile", "Capacitor Android dan plugin widget", "Aplikasi Android, deep link, berbagi PDF, widget belajar"],
            ["Lokalisasi", "Layanan i18n internal", "Antarmuka dan keluaran AI dalam enam bahasa"],
        ],
        widths=[1.15, 2.25, 3.4],
    )

    add_subheading(doc, "Model Data Utama")
    add_table(
        doc,
        ["Entitas", "Isi", "Perlindungan atau Relasi"],
        [
            ["Notebook", "Judul, ringkasan, transkripsi, sumber, flashcard", "Milik pengguna dan dapat ditempatkan pada material"],
            ["Material", "Nama, induk, urutan", "Self reference untuk hierarki; penghapusan tidak menghapus notebook"],
            ["Material Relationship", "Sumber, target, tipe, origin, confidence", "Relasi unik per pengguna dan tipe"],
            ["Quiz", "Soal JSON, kesulitan, fokus, progres", "Cakupan notebook atau mata pelajaran"],
            ["Chat", "Peran, isi, metadata sumber", "Riwayat per notebook, mata pelajaran, atau global"],
            ["Study Day", "Tanggal lokal pengguna", "Satu catatan per pengguna per hari"],
        ],
        widths=[1.55, 2.45, 2.8],
    )
    add_body(
        doc,
        "Kebijakan Row Level Security pada tabel migrasi membatasi kuis, material, relasi, percakapan, dan hari belajar "
        "berdasarkan identitas pengguna. Operasi layanan juga menambahkan filter user_id pada pembacaan dan perubahan data.",
    )

    page_break(doc)

    # Demo guide and judging criteria
    add_section_title(doc, "Skenario Demonstrasi untuk Juri")
    add_body(
        doc,
        "Skenario berikut menunjukkan alur paling representatif dalam sekitar lima menit. Gunakan satu foto catatan dan "
        "satu berkas teks agar proses analisis stabil selama demonstrasi.",
    )
    add_table(
        doc,
        ["Waktu", "Tindakan Demonstrasi", "Hal yang Perlu Diamati"],
        [
            ["0.00 sampai 0.40", "Masuk dan buka dashboard", "Statistik notebook, mata pelajaran, interaksi, dan streak"],
            ["0.40 sampai 1.30", "Unggah foto serta TXT lalu proses", "OCR, judul, klasifikasi, rangkuman, dan tahap review"],
            ["1.30 sampai 2.05", "Buka notebook dan klik sitasi", "Kutipan sumber atau area gambar dapat diperiksa"],
            ["2.05 sampai 2.55", "Buka chat dan ubah pilihan konteks", "Jawaban menggunakan notebook yang dipilih dan sitasi tetap konsisten"],
            ["2.55 sampai 3.45", "Buat kuis atau flashcard", "Kesulitan, fokus, progres, penilaian, dan penjelasan jawaban"],
            ["3.45 sampai 4.30", "Buka graf materi dan relasi", "Hubungan manual serta saran AI antar notebook"],
            ["4.30 sampai 5.00", "Ekspor panduan atau tunjukkan widget", "Hasil belajar dapat dibagikan dan kebiasaan harian dipantau"],
        ],
        widths=[1.25, 2.65, 2.9],
    )

    add_subheading(doc, "Poin Penilaian yang Dapat Diverifikasi")
    add_bullet(doc, "Relevansi masalah terlihat dari dukungan terhadap catatan sehari-hari dan kebutuhan verifikasi jawaban AI.")
    add_bullet(doc, "Kedalaman teknis terlihat pada alur multimodal, sitasi terstruktur, penyimpanan cloud, graf, dan integrasi Android.")
    add_bullet(doc, "Kontrol pengguna terlihat pada tahap review, pemilihan konteks, pengelolaan hierarki, dan relasi manual.")
    add_bullet(doc, "Dampak pembelajaran terlihat pada perpindahan dari rangkuman menuju kuis, flashcard, progres, dan streak.")
    add_bullet(doc, "Skalabilitas konsep terlihat pada cakupan notebook, mata pelajaran, dan seluruh basis pengetahuan pengguna.")

    add_section_title(doc, "Status Kesiapan dan Batas Implementasi")
    add_body(
        doc,
        "Bagian ini menjelaskan kondisi repository saat dokumen disusun. Informasi ini membantu tim memilih jalur demo "
        "yang stabil dan menghindari klaim yang belum sepenuhnya terhubung pada runtime.",
    )
    add_table(
        doc,
        ["Temuan", "Dampak", "Tindakan Sebelum Penjurian"],
        [
            ["PDF telah diekstrak per halaman, tetapi cabang analisis AI saat ini hanya meneruskan tipe text dan image", "Isi PDF berpotensi tidak masuk ke analisis notebook", "Normalisasi sumber PDF menjadi text atau tambahkan cabang pdf pada payload AI"],
            ["Pengaturan provider dan model tersimpan di profil, tetapi saveLLM belum menyinkronkan nilai tersebut ke aiService", "Perubahan pada tab LLM belum tentu memengaruhi runtime", "Hubungkan base URL, API key, dan model ke konfigurasi layanan saat penyimpanan"],
            ["Tab History pada panel konteks masih berupa placeholder", "Riwayat pesan tersimpan, tetapi belum ada daftar percakapan terpisah", "Demonstrasikan riwayat per cakupan; jangan mengklaim pengelolaan multi percakapan"],
            ["Berkas sumber memakai public URL dari Supabase Storage", "Materi sensitif memerlukan perlindungan tambahan", "Gunakan bucket privat dan signed URL untuk penggunaan produksi"],
            ["Belum terlihat rangkaian pengujian aplikasi otomatis pada repository", "Risiko regresi perlu ditangani lewat uji manual", "Jalankan build dan checklist demo pada web serta Android sebelum presentasi"],
        ],
        widths=[2.65, 1.8, 2.35],
    )

    page_break(doc)

    # Video script
    add_section_title(doc, "Naskah Video Penjelasan Karya")
    add_body(
        doc,
        "Durasi sasaran dua menit tiga puluh detik. Narasi dibuat untuk video layar dengan potongan singkat pada setiap "
        "fitur utama.",
    )
    add_table(
        doc,
        ["Durasi", "Visual", "Narasi"],
        [
            ["0.00 sampai 0.15", "Logo dan dashboard", "Ini adalah Ibuki Study, ruang belajar berbasis AI yang mengubah catatan menjadi pengetahuan terstruktur, dapat diperiksa, dan siap dipakai untuk latihan."],
            ["0.15 sampai 0.40", "Unggah foto dan TXT", "Pelajar biasanya memiliki foto papan tulis, dokumen, dan catatan teks yang tersebar. Di Ibuki Study, sumber tersebut dikumpulkan dalam satu notebook dan diproses menjadi judul, transkripsi, rangkuman, serta jalur materi."],
            ["0.40 sampai 1.05", "Klik sitasi dan buka sumber", "Setiap fakta yang berasal dari materi membawa sitasi. Pengguna dapat membuka kutipan, halaman PDF, atau area tertentu pada gambar. Dengan demikian, jawaban AI dapat diperiksa tanpa meninggalkan aplikasi."],
            ["1.05 sampai 1.30", "Panel konteks pada chat", "Chat tersedia untuk satu notebook, satu mata pelajaran, atau seluruh koleksi. Pengguna memilih sendiri sumber yang dipakai, dan aplikasi menyimpan konteks sitasi setiap jawaban agar referensi lama tetap benar."],
            ["1.30 sampai 1.55", "Pohon dan graf materi", "Materi diatur melalui hierarki yang fleksibel dan graf hubungan. AI dapat menyarankan keterkaitan, sementara pengguna tetap dapat membuat atau memperbaiki relasi secara manual."],
            ["1.55 sampai 2.15", "Kuis, hasil, flashcard", "Rangkuman kemudian menjadi aktivitas belajar. Ibuki Study membuat kuis sesuai tingkat kesulitan dan fokus, menyimpan progres serta nilai terbaik, dan menyediakan flashcard untuk pengulangan."],
            ["2.15 sampai 2.30", "Ekspor PDF dan widget Android", "Hasil percakapan dapat disimpan sebagai materi atau dibagikan sebagai PDF. Pada Android, widget menampilkan streak dan akses cepat untuk belajar. Ibuki Study menyatukan sumber, penjelasan, verifikasi, dan latihan dalam satu alur."],
        ],
        widths=[1.15, 1.85, 3.8],
    )

    add_subheading(doc, "Daftar Pengambilan Gambar")
    add_bullet(doc, "Rekam layar dengan data contoh yang singkat dan mudah dibaca.")
    add_bullet(doc, "Perbesar kursor saat membuka sitasi agar mekanisme verifikasi terlihat jelas.")
    add_bullet(doc, "Gunakan satu transisi sederhana antar layar dan hindari menampilkan API key atau data akun.")
    add_bullet(doc, "Tampilkan hasil kuis dan graf cukup lama agar juri dapat membaca status serta tipe relasinya.")

    add_section_title(doc, "Jejak Implementasi")
    add_body(
        doc,
        "Komponen berikut menjadi rujukan teknis utama untuk memeriksa kesesuaian penjelasan dengan repository.",
    )
    add_table(
        doc,
        ["Area", "Berkas Utama"],
        [
            ["Navigasi dan cakupan aplikasi", "src/App.vue"],
            ["Pembuatan notebook", "src/components/CreateNotebookView.vue dan src/services/ai.js"],
            ["Sitasi dan penampil sumber", "src/services/citations.js dan src/components/SourceViewerModal.vue"],
            ["Chat dan pemilih konteks", "src/components/ChatView.vue dan ChatContextSidebar.vue"],
            ["Kuis, flashcard, dan streak", "QuizPanel.vue, FlashcardView.vue, db.js, dan widget.js"],
            ["Pohon dan graf pengetahuan", "MaterialManager.vue, materials.js, dan MaterialGraphView.vue"],
            ["Data dan keamanan", "src/services/db.js serta migrations"],
            ["Android", "capacitor.config.json dan android/app/src/main"],
        ],
        widths=[2.35, 4.45],
    )

    doc.core_properties.title = "Penjelasan Karya Ibuki Study"
    doc.core_properties.subject = "Acuan penilaian juri mengenai ide, cara kerja, dan fitur unik"
    doc.core_properties.author = "Tim Ibuki Study"
    doc.core_properties.keywords = "Ibuki Study, penjelasan karya, juri, AI, pembelajaran"
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
